import hashlib
import io
import json
from datetime import date, timedelta

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import func, select

from superflow.models import Document, DocumentChunk, DocumentVersion, QuestionRun
from superflow.parsing import MAX_BYTES, ParseError, parse_file
from superflow.seed import seed


def metadata(**overrides):
    return {
        "title": "Synthetic Training Guide",
        "domain": "People",
        "owner_id": "sme",
        "authority_level": 4,
        "access_roles": ["employee", "sme", "governance"],
        "related_document_ids": [],
        **overrides,
    }


def upload(client, headers, content=b"# Training\n\nTraining starts at nine each Monday.", **extra):
    return client.post(
        "/api/documents",
        headers=headers,
        data={
            "metadata": json.dumps(metadata()),
            "version": "1.0",
            "effective_date": "2026-01-01",
            **extra,
        },
        files={"file": ("training.md", content, "text/markdown")},
    )


def ask(client, headers, question="How many annual leave days do employees receive?", **filters):
    return client.post("/api/ask", headers=headers, json={"question": question, **filters})


def test_authentication_required(client):
    assert client.get("/api/documents").status_code == 401
    assert client.get("/api/documents", headers={"X-User-Id": "unknown"}).status_code == 401
    assert len(client.get("/api/users").json()) == 6


def test_role_inventory_and_direct_read_non_leakage(client, employee, admin):
    allowed = client.get("/api/documents", headers=employee).json()
    assert "proposal-a" not in {document["id"] for document in allowed}
    assert "regulation" not in {document["id"] for document in allowed}
    assert client.get("/api/documents/proposal-a", headers=employee).status_code == 404
    missing = client.get("/api/documents/unknown", headers=employee)
    assert missing.json() == client.get("/api/documents/proposal-a", headers=employee).json()
    assert len(client.get("/api/documents", headers=admin).json()) == 12


def test_retrieval_non_leakage_even_with_restricted_filter(client, employee):
    result = ask(
        client, employee, "What is the Proposal A initial investment?", document_ids=["proposal-a"]
    ).json()
    assert result["outcome"] == "no_sources"
    assert result["sources"] == []
    assert "180,000,000" not in result["answer"]


def test_active_version_and_exact_citations(client, employee):
    response = ask(client, employee, document_ids=["leave-policy"])
    assert response.status_code == 201
    result = response.json()
    assert result["outcome"] == "source_grounded"
    assert "15 days" in result["answer"]
    assert "12 days" not in result["answer"]
    assert all(
        source["version"] == "2.0" and source["status"] == "active" for source in result["sources"]
    )
    for source in result["sources"]:
        assert source["excerpt"] in result["answer"]
        assert f"[{source['citation_id']}]" in result["answer"]
        assert source["owner"] and source["effective_date"] and source["section_path"]


def test_historical_versions_are_explicit_and_labeled(client, employee):
    result = ask(client, employee, document_ids=["leave-policy"], include_outdated=True).json()
    assert any(
        source["version"] == "1.0" and source["status"] == "outdated"
        for source in result["sources"]
    )


def test_domain_filters_and_no_evidence(client, employee):
    result = ask(client, employee, domain="Finance").json()
    assert result["outcome"] == "no_sources"
    result = ask(client, employee, "quasar neutrino spectrum").json()
    assert result["sources"] == []
    assert "could not find supporting evidence" in result["answer"]
    assert "Verified" not in result["answer"]


def test_concept_retrieval(client, employee):
    result = ask(client, employee, "vacation entitlement", document_ids=["leave-policy"]).json()
    assert result["sources"]
    assert "15 days" in result["answer"]


def test_saved_question_owner_and_reauthorized_sources(client, admin, employee, db):
    result = ask(
        client, {"X-User-Id": "finance"}, "What does Proposal A cost?", document_ids=["proposal-a"]
    ).json()
    run_id = result["id"]
    assert client.get(f"/api/ask/{run_id}", headers=employee).status_code == 404
    assert client.get(f"/api/ask/{run_id}/sources", headers=employee).status_code == 404
    document = db.get(Document, "proposal-a")
    document.access_roles = ["governance"]
    db.commit()
    assert client.get(f"/api/ask/{run_id}", headers={"X-User-Id": "finance"}).status_code == 403
    history = client.get("/api/ask", headers={"X-User-Id": "finance"}).json()
    assert run_id not in {run["id"] for run in history}


def test_upload_persistence_checksum_and_structure(client, admin, db):
    response = upload(client, admin)
    assert response.status_code == 201, response.text
    document = response.json()
    assert document["versions"][0]["chunks"][0]["section_path"] == "Training"
    assert (
        document["versions"][0]["checksum"]
        == hashlib.sha256(b"# Training\n\nTraining starts at nine each Monday.").hexdigest()
    )
    assert db.get(Document, document["id"]) is not None
    result = ask(
        client,
        {"X-User-Id": "employee"},
        "When does training start?",
        document_ids=[document["id"]],
    ).json()
    assert "nine" in result["answer"]


def test_write_permissions(client, employee):
    assert upload(client, employee).status_code == 403
    assert (
        client.patch("/api/documents/leave-policy", headers=employee, json=metadata()).status_code
        == 403
    )
    assert client.post("/api/documents/leave-policy/index", headers=employee).status_code == 403


def test_version_activation_supersedes_and_duplicate_rolls_back(client, admin):
    response = client.post(
        "/api/documents/leave-policy/versions",
        headers=admin,
        data={"version": "3.0", "effective_date": "2026-02-01"},
        files={"file": ("leave.md", b"# Leave\nEmployees receive 20 days of annual leave.")},
    )
    assert response.status_code == 201, response.text
    versions = response.json()["versions"]
    assert sum(version["status"] == "active" for version in versions) == 1
    assert (
        next(version for version in versions if version["version"] == "2.0")["status"] == "outdated"
    )
    result = ask(client, admin, document_ids=["leave-policy"]).json()
    assert "20 days" in result["answer"] and "15 days" not in result["answer"]
    duplicate = client.post(
        "/api/documents/leave-policy/versions",
        headers=admin,
        data={"version": "3.0", "effective_date": "2026-02-01"},
        files={"file": ("leave.md", b"Duplicate content")},
    )
    assert duplicate.status_code == 409


def test_draft_excluded_then_published(client, admin):
    document = upload(client, admin, status="draft").json()
    assert (
        ask(client, admin, "Training Monday", document_ids=[document["id"]]).json()["sources"] == []
    )
    version_id = document["versions"][0]["id"]
    response = client.patch(
        f"/api/documents/{document['id']}/versions/{version_id}?status=active", headers=admin
    )
    assert response.status_code == 200
    assert ask(client, admin, "Training Monday", document_ids=[document["id"]]).json()["sources"]


def test_future_active_version_rejected(client, admin):
    future = (date.today() + timedelta(days=1)).isoformat()
    assert upload(client, admin, effective_date=future).status_code == 422


def test_metadata_edit_and_related_permission_validation(client, admin, employee):
    response = client.patch(
        "/api/documents/leave-policy", headers=admin, json=metadata(title="Updated Leave Policy")
    )
    assert response.status_code == 200
    assert (
        client.get("/api/documents/leave-policy", headers=employee).json()["title"]
        == "Updated Leave Policy"
    )
    response = client.patch(
        "/api/documents/leave-policy",
        headers={"X-User-Id": "sme"},
        json=metadata(related_document_ids=["proposal-a"]),
    )
    assert response.status_code == 404


@pytest.mark.parametrize(
    "changes",
    [
        {"authority_level": 9},
        {"access_roles": ["unknown"]},
        {"title": "   "},
        {"owner_id": "unknown"},
        {"domain": "Unknown"},
        {"access_roles": []},
    ],
)
def test_invalid_metadata(client, admin, changes):
    response = client.post(
        "/api/documents",
        headers=admin,
        data={"metadata": json.dumps(metadata(**changes))},
        files={"file": ("file.txt", b"Some text")},
    )
    assert response.status_code == 422


def test_reindex_preserves_citation_ids_and_download_permission(client, admin, employee):
    result = ask(
        client, admin, document_ids=["proposal-a"], question="Proposal A investment"
    ).json()
    ids = [source["chunk_id"] for source in result["sources"]]
    assert client.post("/api/documents/proposal-a/index", headers=admin).status_code == 200
    saved = client.get(f"/api/ask/{result['id']}", headers=admin).json()
    assert [source["chunk_id"] for source in saved["sources"]] == ids
    version_id = result["sources"][0]["version_id"]
    route = f"/api/documents/proposal-a/versions/{version_id}/file"
    assert client.get(route, headers=employee).status_code == 404
    assert b"180,000,000" in client.get(route, headers=admin).content


def test_seed_is_idempotent(db):
    counts_before = [
        db.scalar(select(func.count()).select_from(model))
        for model in [Document, DocumentVersion, DocumentChunk, QuestionRun]
    ]
    assert seed(db) == 0
    counts_after = [
        db.scalar(select(func.count()).select_from(model))
        for model in [Document, DocumentVersion, DocumentChunk, QuestionRun]
    ]
    assert counts_before == counts_after


def test_txt_chunks_and_nul_sanitization():
    _, chunks = parse_file("sample.txt", ("hello\x00 " + "longtext " * 800).encode())
    assert len(chunks) > 1
    assert all(len(chunk.text) <= 1200 and "\x00" not in chunk.text for chunk in chunks)


def test_pdf_parsing_page_citations():
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(50, 750, "Synthetic policy requires manager approval.")
    canvas.showPage()
    canvas.save()
    _, chunks = parse_file("policy.pdf", buffer.getvalue())
    assert chunks[0].section == "Page 1"
    assert "manager approval" in chunks[0].text


def test_docx_parsing_headings_and_tables():
    document = DocxDocument()
    document.add_heading("Approval steps", 1)
    document.add_paragraph("Obtain manager approval.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Owner", "Finance"
    buffer = io.BytesIO()
    document.save(buffer)
    _, chunks = parse_file("policy.docx", buffer.getvalue())
    assert chunks[0].section == "Approval steps"
    assert any(chunk.section == "Table 1" and "Finance" in chunk.text for chunk in chunks)


def test_safe_parser_failures():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ParseError, match="OCR"):
        parse_file("scan.pdf", buffer.getvalue())
    with pytest.raises(ParseError):
        parse_file("broken.docx", b"not a zip")
    with pytest.raises(ParseError):
        parse_file("code.exe", b"not supported")
    with pytest.raises(ParseError, match="10 MB"):
        parse_file("large.txt", b"a" * (MAX_BYTES + 1))


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_binary_document_upload_end_to_end(client, admin, format):
    buffer = io.BytesIO()
    if format == "pdf":
        canvas = Canvas(buffer)
        canvas.drawString(50, 750, "Training requires manager approval every Monday.")
        canvas.save()
    else:
        document = DocxDocument()
        document.add_heading("Training", 1)
        document.add_paragraph("Training requires manager approval every Monday.")
        document.save(buffer)
    response = client.post(
        "/api/documents",
        headers=admin,
        data={"metadata": json.dumps(metadata()), "effective_date": "2026-01-01"},
        files={"file": (f"training.{format}", buffer.getvalue())},
    )
    assert response.status_code == 201, response.text
    result = ask(
        client, admin, "Training manager approval", document_ids=[response.json()["id"]]
    ).json()
    assert "every Monday" in result["answer"]
    assert result["sources"][0]["section_path"] == ("Page 1" if format == "pdf" else "Training")


def test_failed_upload_and_version_are_atomic(client, admin, db):
    before = db.scalar(select(func.count()).select_from(Document))
    assert upload(client, admin, content=b"").status_code == 422
    assert db.scalar(select(func.count()).select_from(Document)) == before
    response = client.post(
        "/api/documents/leave-policy/versions",
        headers=admin,
        data={"version": "bad", "effective_date": "2026-01-01"},
        files={"file": ("empty.md", b"")},
    )
    assert response.status_code == 422
    document = client.get("/api/documents/leave-policy", headers=admin).json()
    assert len(document["versions"]) == 2
    assert (
        next(version for version in document["versions"] if version["status"] == "active")[
            "version"
        ]
        == "2.0"
    )


def test_later_phase_routes_are_absent(client, admin):
    for endpoint in [
        "cases",
        "approvals",
        "audit",
        "analytics/overview",
        "regulatory/compare",
        "finance/calculate",
    ]:
        assert client.get(f"/api/{endpoint}", headers=admin).status_code == 404
